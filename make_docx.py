import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

# Main Title
title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(12)
title.paragraph_format.space_after = Pt(2)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
t_run = title.add_run('KỊCH BẢN & LỜI THOẠI THUYẾT TRÌNH')
t_run.font.size = Pt(18)
t_run.font.bold = True
t_run.font.color.rgb = RGBColor(37, 99, 235)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
s_run = subtitle.add_run('HỌC THUYẾT HÌNH THÁI KINH TẾ - XÃ HỘI (TRIẾT HỌC MÁC - LÊNIN)\nChuẩn 25 Slide Thuyết Trình • Kịch bản phân chia 4 thành viên')
s_run.font.size = Pt(11)
s_run.font.italic = True
s_run.font.color.rgb = RGBColor(100, 116, 139)

with open('/home/amtia/mln111-ass/kich_ban_thuyet_trinh.md', 'r', encoding='utf-8') as f:
    lines = f.readlines()

for line in lines:
    line = line.strip()
    if not line:
        continue
    if line.startswith('# KỊCH BẢN') or line.startswith('## HỌC THUYẾT') or line.startswith('*Website') or line == '---':
        continue
    elif line.startswith('## '):
        h = doc.add_heading(line.replace('## ', ''), level=1)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        for r in h.runs:
            r.font.size = Pt(13.5)
            r.font.color.rgb = RGBColor(15, 23, 42)
    elif line.startswith('### 📌 '):
        h = doc.add_heading(line.replace('### 📌 ', ''), level=2)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(3)
        for r in h.runs:
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(2, 132, 199)
    elif line.startswith('> *'):
        quote_text = line.replace('> *', '').replace('*', '').strip()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(5)
        r_tag = p.add_run('🎙️ LỜI THOẠI: ')
        r_tag.bold = True
        r_tag.font.color.rgb = RGBColor(30, 64, 175)
        r_txt = p.add_run(quote_text)
        r_txt.font.italic = True
        r_txt.font.color.rgb = RGBColor(30, 41, 59)
    elif line.startswith('- **Thao tác:**'):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run('👉 Thao tác người nói: ')
        r.bold = True
        r.font.color.rgb = RGBColor(217, 119, 6)
        p.add_run(line.replace('- **Thao tác:**', '').strip())
    else:
        p = doc.add_paragraph(line)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)

doc.save('/home/amtia/mln111-ass/Kich_Ban_Thuyet_Trinh_MLN111.docx')
print('Successfully saved Kich_Ban_Thuyet_Trinh_MLN111.docx')
